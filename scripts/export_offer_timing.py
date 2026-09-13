#!/usr/bin/env python3
"""Generate active Assessment offer-timing runtime from reviewed Offer Timing.xlsx.

Phase 5C: the workbook is authoritative. Timing math is recomputed from Current Offer
Input + Offer History + Timing Config, while user-facing display labels come from the
reviewed Runtime Display Snapshot sheet. No existing runtime JSON is read as input.
"""
from __future__ import annotations
import argparse, json, math
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
import openpyxl
from openpyxl.utils.datetime import from_excel

ROOT=Path(__file__).resolve().parents[1]

def records(ws, header_row=1):
    headers=list(next(ws.iter_rows(min_row=header_row,max_row=header_row,values_only=True)))
    for vals in ws.iter_rows(min_row=header_row+1,values_only=True):
        if any(v is not None for v in vals): yield dict(zip(headers,vals))

def manifest(wb,key):
    for r in records(wb['Manifest']):
        if r['Field']==key: return r['Value']
    raise KeyError(key)

def as_date(v,epoch):
    if isinstance(v,(int,float)): v=from_excel(v,epoch)
    if isinstance(v,datetime): return v.date()
    if isinstance(v,date): return v
    if isinstance(v,str): return datetime.fromisoformat(v).date()
    raise ValueError(f'Unsupported date {v!r}')

def days_to_quarter(d,q):
    current=(d.month-1)//3+1
    if current==q: return 0
    month=1+(q-1)*3
    year=d.year if month>d.month else d.year+1
    return (date(year,month,1)-d).days

def rating(current, high, pct, cfg):
    ratio=1 if max(current,high)==0 else current/max(current,high)
    if ratio>=1: return '史高'
    if (ratio>=cfg['Excellent Ratio'] or
        (pct>=cfg['Excellent Percentile'] and ratio>=cfg['Excellent Percentile Min Ratio']) or
        ratio>=cfg['Good Ratio'] or
        (ratio>=cfg['Good Combo Ratio'] and pct>=cfg['Good Combo Percentile']) or
        (pct>=cfg['Good Percentile'] and ratio>=cfg['Good Percentile Min Ratio'])): return '较高'
    if (ratio>=cfg['Normal Absolute Ratio Floor'] or
        (ratio>=cfg['Normal Ratio'] and pct>=cfg['Normal Percentile']) or
        (ratio>=cfg['Normal Alt Ratio'] and pct>=cfg['Normal Alt Percentile'])): return '一般'
    return '较低'

def generate(workbook: Path):
    wb=openpyxl.load_workbook(workbook,read_only=True,data_only=True)
    if str(manifest(wb,'status')).upper()!='FINAL': raise ValueError('Offer Timing source is not FINAL')
    snapshot=str(manifest(wb,'snapshot_date'))
    cfg={str(r['Timing Config']):float(r['Value']) for r in records(wb['Timing Config']) if r.get('Timing Config')}
    current={str(r['Offer Timing ID']):r for r in records(wb['Current Offer Input']) if r.get('Offer Timing ID')}
    hist=defaultdict(list)
    for r in records(wb['Offer History']):
        if r.get('Offer Timing ID'): hist[str(r['Offer Timing ID'])].append(r)
    displays=list(records(wb['Runtime Display Snapshot'],header_row=2))
    out={}
    for drow in displays:
        oid=str(drow['Offer Timing ID']); pid=str(drow['Product ID'])
        if oid not in current: raise ValueError(f'Missing current offer {oid}')
        c=current[oid]; cv=float(c['Current Comparable Value']); unit=str(c['Comparison Unit']); ev=as_date(c['Evaluation Date'],wb.epoch)
        rows=[r for r in hist[oid] if str(r.get('Timing Eligible')).upper()=='YES' and r.get('Comparison Unit')==unit and isinstance(r.get('Comparable Value'),(int,float))]
        high=max([float(r['Comparable Value']) for r in rows],default=cv); effective_high=max(cv,high)
        pct=1 if not rows else sum(1 for r in rows if float(r['Comparable Value'])<=cv)/len(rows)
        rat=rating(cv,effective_high,pct,cfg)
        qyears=defaultdict(set)
        for r in rows:
            if float(r['Comparable Value'])<=cv: continue
            if int(r.get('Year-Quarter First Eligible') or 0)!=1: continue
            q=int(r.get('Quarter') or 0); y=int(r.get('Year') or 0)
            if q and y: qyears[q].add(y)
        eligible_q=[q for q,years in qyears.items() if len(years)>=int(cfg['Future Timing Min Years'])]
        days=min([days_to_quarter(ev,q) for q in eligible_q],default=None)
        wait='建议等待' if rat!='史高' and days is not None and days<=cfg['Future Timing Wait Window Days'] else '不建议等待'
        result={'rating':rat,'wait':wait,'historical_high_label':str(drow['Historical High Display']),'historical_low_label':str(drow['Historical Low Display']),'history_count':int(drow['History Display Count'])}
        if drow.get('Maximum Offer Display'): result['maximum_offer_label']=str(drow['Maximum Offer Display'])
        out[pid]={'offer_timing_id':oid,'product_id':pid,'snapshot_date':snapshot,'current_offer':{'offer_label':str(drow['Current Offer Display']),'offer_mechanism':str(c['Offer Mechanism']),'comparable_value':c['Current Comparable Value'],'comparison_unit':unit,'expiry':drow.get('Expiry')},'timing_result':result}
    return out

def main():
    p=argparse.ArgumentParser(); p.add_argument('workbook',type=Path); p.add_argument('--output',type=Path,default=ROOT/'data/offer-timing.json'); a=p.parse_args()
    out=generate(a.workbook); a.output.parent.mkdir(parents=True,exist_ok=True); a.output.write_text(json.dumps(out,ensure_ascii=False,indent=2,default=str)+'\n',encoding='utf-8')
    print(json.dumps({'products':len(out),'output':str(a.output),'status':'PASS'},ensure_ascii=False))
if __name__=='__main__': main()
