#!/usr/bin/env python3
"""Export deterministic machine-value acceptance cases from reviewed QA workbooks.

This is an offline publication boundary. It never invents answers: common/specific
question labels are resolved through data/questions.json, benefit selections come from
the workbook's Benefit Persona Map IDs, and expected outputs come from the reviewed
Assessment Acceptance register.
"""
from __future__ import annotations
import argparse, hashlib, json, re
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]

def load_json(path: Path):
    return json.loads(path.read_text(encoding='utf-8'))

def sha256(path: Path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

def rows_as_dicts(ws):
    headers=[c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        yield dict(zip(headers,row))

def option_map(block: dict) -> dict:
    return {str(o['label']):o['value'] for o in (block.get('options') or [])}

def flatten_blocks(card_questions: dict):
    blocks={}; number_ids=set()
    def add(block):
        if block.get('type') in {'single','multi'} and block.get('id'):
            blocks[block['id']]=block
        for sub in block.get('subs') or []:
            blocks[sub['id']]=sub
            if sub.get('number_id'):
                number_ids.add(sub['number_id'])
    add(card_questions['q5']); add(card_questions['q6'])
    return blocks, number_ids

def resolve_option(block: dict, raw):
    m=option_map(block)
    if block.get('type')=='multi':
        if raw is None:
            return []
        text=str(raw)
        if text in m:
            return [m[text]]
        parts=[x.strip() for x in re.split(r'[;；]', text) if x.strip()]
        missing=[x for x in parts if x not in m]
        if missing:
            raise ValueError(f"unknown multi option label(s) for {block.get('id')}: {missing}")
        return [m[x] for x in parts]
    key=str(raw)
    if key not in m:
        raise ValueError(f"unknown option label for {block.get('id')}: {raw!r}")
    return m[key]

def split_ids(raw):
    if raw is None:
        return []
    parts=[x.strip() for x in str(raw).split(';') if x.strip()]
    return [] if parts==['NONE'] else parts

def split_rules(raw):
    if not raw:
        return []
    return [x.strip() for x in re.split(r'[,;；]', str(raw)) if x.strip()]

def main():
    ap=argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--e2e', type=Path, required=True)
    ap.add_argument('--acceptance', type=Path, required=True)
    ap.add_argument('--questions', type=Path, default=ROOT/'data/questions.json')
    ap.add_argument('--products', type=Path, default=ROOT/'data/products.json')
    ap.add_argument('--registry', type=Path, default=ROOT/'config/qa-source-registry.json')
    ap.add_argument('--output-dir', type=Path, default=ROOT/'data/qa')
    ap.add_argument('--manifest-output', type=Path, default=ROOT/'data/qa/fixed-profile-acceptance-manifest.json')
    args=ap.parse_args()

    registry=load_json(args.registry)
    sources=registry['sources']
    checks={'assessment_e2e_fixed_profiles':args.e2e,'assessment_acceptance':args.acceptance}
    for key,path in checks.items():
        actual=sha256(path); expected=sources[key]['sha256']
        if actual!=expected:
            raise SystemExit(f"QA source SHA256 changed for {key}: {actual}; reviewed lock is {expected}")

    questions=load_json(args.questions); products=load_json(args.products)
    common={q['id']:q for q in questions['common']}
    product_by_name={p['name']:pid for pid,p in products['cards'].items()}
    e2e=load_workbook(args.e2e, data_only=True, read_only=True)
    acceptance=load_workbook(args.acceptance, data_only=True, read_only=True)

    evaluation_date=None
    for row in e2e['System Facts'].iter_rows(min_row=2, values_only=True):
        if row[0]=='Evaluation Date':
            evaluation_date=str(row[1]); break
    if not evaluation_date:
        raise SystemExit('missing frozen Evaluation Date in System Facts')

    benefit_map={}
    for d in rows_as_dicts(e2e['Benefit Persona Map']):
        benefit_map[(d['Persona'],d['Card'])]={'q7':split_ids(d['Q7 Benefit IDs']),'q8':split_ids(d['Q8 Benefit IDs'])}

    expected={}
    for d in rows_as_dicts(acceptance['Fixed Profile 120 Register']):
        expected[d['Run ID']]={
          'applicationImpact':d['Application Impact'],
          'applicationSoftImpact':d['Application Soft Impact'],
          'applicationHardBehavior':d['Application Hard Behavior'],
          'bonusImpact':d['Bonus Impact'],
          'bonusHardBehavior':d['Bonus Hard Behavior'],
          'baseApproval':d['Base Approval'],
          'approvalAfterSoftRules':d['Approval After Soft Rules'],
          'adjustedApproval':d['Adjusted Approval'],
          'approvalLabel':d['获批可能性'],
          'offerRating':d['开卡奖励评级'],
          'waitRecommendation':d['Wait Recommendation'],
          'bonusLabel':d['能否拿奖励'],
          'longTermLabel':d['长期持有价值'],
          'recommendedAction':d['Recommended Action'],
          'systemMissingNotes':d['System/Missing Notes'],
          'benefitPersona':d['Benefit Persona'],
        }

    cases=[]
    for d in rows_as_dicts(e2e['120 E2E']):
        run_id=d['Run ID']; card_name=d['Card']
        if card_name not in product_by_name:
            raise SystemExit(f"unknown card in E2E workbook: {card_name}")
        pid=product_by_name[card_name]
        blocks, number_ids=flatten_blocks(questions['cards'][pid])
        raw_specific=json.loads(d['Q5/Q6 Answers'])
        specific={}
        for qid,raw in raw_specific.items():
            if isinstance(raw,str) and raw.startswith('未显示（'):
                continue
            if qid in number_ids:
                specific[qid]=int(raw); continue
            if qid not in blocks:
                raise SystemExit(f"unknown Question ID in {run_id}: {qid}")
            specific[qid]=resolve_option(blocks[qid],raw)
        b=benefit_map[(d['Benefit Persona'],card_name)]
        answers={
          'a1':resolve_option(common['a1'],d['A1']),
          'a2':resolve_option(common['a2'],d['A2']),
          'a3':resolve_option(common['a3'],d['A3']),
          'a4':resolve_option(common['a4'],d['A4']),
          'specific':specific,
          'q7':b['q7'],'q8':b['q8'],'q7Answered':True,'q8Answered':True,
        }
        ex=dict(expected[run_id]); ex['internalState']=d['Internal State']; ex['triggeredRules']=split_rules(d['Triggered Rule IDs'])
        cases.append({'run_id':run_id,'profile':d['Profile'],'product_id':pid,'card_code':d['Card Code'],'evaluation_date':evaluation_date,'answers':answers,'expected':ex})

    if len(cases)!=registry['fixture']['case_count']:
        raise SystemExit(f"case count mismatch: {len(cases)}")
    if evaluation_date!=registry['fixture']['evaluation_date']:
        raise SystemExit(f"evaluation date mismatch: {evaluation_date}")
    import gzip
    raw_payload={'schema_version':registry['fixture']['schema_version'],'fixture_id':registry['fixture']['fixture_id'],'evaluation_date':evaluation_date,'source_provenance':{k:{'file_name':sources[k]['file_name'],'drive_id':sources[k]['drive_id'],'sha256':sources[k]['sha256']} for k in ('assessment_e2e_fixed_profiles','assessment_acceptance')},'case_count':len(cases),'cases':cases}
    raw=(json.dumps(raw_payload,ensure_ascii=False,separators=(',',':'))+'\n').encode('utf-8')
    artifact_path=args.output_dir/'fixed-profile-acceptance.json.gz'
    artifact_path.parent.mkdir(parents=True,exist_ok=True)
    compressed=gzip.compress(raw,compresslevel=9,mtime=0);artifact_path.write_bytes(compressed)
    git_blob=hashlib.sha1(b'blob '+str(len(compressed)).encode()+b'\0'+compressed).hexdigest()
    manifest={'schema_version':registry['fixture']['schema_version'],'fixture_id':registry['fixture']['fixture_id'],'evaluation_date':evaluation_date,'case_count':len(cases),'source_provenance':raw_payload['source_provenance'],'artifact':{'path':str(artifact_path.relative_to(ROOT)).replace('\\','/'),'encoding':'gzip-json','git_blob_sha':git_blob,'sha256':hashlib.sha256(compressed).hexdigest(),'uncompressed_sha256':hashlib.sha256(raw).hexdigest(),'uncompressed_bytes':len(raw),'compressed_bytes':len(compressed)}}
    args.manifest_output.parent.mkdir(parents=True,exist_ok=True);args.manifest_output.write_text(json.dumps(manifest,ensure_ascii=False,separators=(',',':'))+'\n',encoding='utf-8')
    print(json.dumps({'status':'PASS','cases':len(cases),'evaluation_date':evaluation_date,'manifest':str(args.manifest_output),'artifact':str(artifact_path)},ensure_ascii=False))

if __name__=='__main__': main()
